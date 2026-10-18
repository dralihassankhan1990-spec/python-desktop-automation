"""
make_sample_data.py — builds project-data.xlsx with FICTIONAL data for two projects.
Run once:  python make_sample_data.py
Re-run any time to reset the sample data.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="0E4C81")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)

PROJECTS = [
    # id, name, manager, sponsor, start, end, budget, spent
    ("ERP-01", "Riverdale ERP Rollout", "Ali Raza", "Operations Director",
     "2026-01-12", "2026-12-15", 12_500_000, 7_800_000),
    ("CRM-02", "Harbor CRM Upgrade", "Fatima Khan", "Head of Sales",
     "2026-03-02", "2026-11-30", 4_200_000, 3_900_000),
]

MILESTONES = [
    # project, milestone, due, status, owner
    ("ERP-01", "Requirements & BRD sign-off", "2026-03-31", "Done", "Ali Raza"),
    ("ERP-01", "Core modules configured (Finance, Inventory)", "2026-06-30", "Done", "Bilal Ahmed"),
    ("ERP-01", "Data migration — master data loaded", "2026-08-15", "Done", "Sana Iqbal"),
    ("ERP-01", "UAT — Round 1 (Finance + Warehouse)", "2026-10-20", "In Progress", "Ali Raza"),
    ("ERP-01", "UAT — Round 2 (all modules)", "2026-11-20", "Not Started", "Fatima Khan"),
    ("ERP-01", "Go-live & hypercare", "2026-12-15", "Not Started", "Ali Raza"),
    ("CRM-02", "CRM selection & vendor contract", "2026-04-15", "Done", "Fatima Khan"),
    ("CRM-02", "Data cleansing & import (12,400 contacts)", "2026-06-30", "Done", "Usman Tariq"),
    ("CRM-02", "Sales pipeline workflows configured", "2026-08-31", "Done", "Fatima Khan"),
    ("CRM-02", "User training — 3 batches", "2026-10-10", "Delayed", "Usman Tariq"),
    ("CRM-02", "Go-live & legacy CRM switch-off", "2026-11-30", "Not Started", "Fatima Khan"),
]

TASKS = [
    # project, task, owner, status, due
    ("ERP-01", "Finalize chart of accounts mapping", "Bilal Ahmed", "Done", "2026-05-15"),
    ("ERP-01", "Configure approval workflows", "Bilal Ahmed", "Done", "2026-06-10"),
    ("ERP-01", "Migrate vendor master (842 records)", "Sana Iqbal", "Done", "2026-07-20"),
    ("ERP-01", "Migrate item master (3,150 SKUs)", "Sana Iqbal", "Done", "2026-08-05"),
    ("ERP-01", "Write UAT test scripts — Finance", "Ali Raza", "Done", "2026-09-25"),
    ("ERP-01", "Write UAT test scripts — Warehouse", "Ali Raza", "In Progress", "2026-10-18"),
    ("ERP-01", "Execute UAT Round 1 — Finance", "Ali Raza", "In Progress", "2026-10-20"),
    ("ERP-01", "Fix UAT defects (batch 1)", "Bilal Ahmed", "Not Started", "2026-10-28"),
    ("ERP-01", "Prepare cutover runbook", "Sana Iqbal", "Not Started", "2026-11-10"),
    ("ERP-01", "Train 45 end users", "Fatima Khan", "Not Started", "2026-11-25"),
    ("CRM-02", "Map legacy fields to new CRM", "Usman Tariq", "Done", "2026-05-20"),
    ("CRM-02", "Import contacts & de-duplicate", "Usman Tariq", "Done", "2026-06-25"),
    ("CRM-02", "Configure dashboards for sales leads", "Fatima Khan", "Done", "2026-08-20"),
    ("CRM-02", "Batch 1 training (25 users)", "Fatima Khan", "Done", "2026-09-15"),
    ("CRM-02", "Batch 2 training (25 users)", "Fatima Khan", "Delayed", "2026-10-10"),
    ("CRM-02", "Batch 3 training (20 users)", "Usman Tariq", "Not Started", "2026-10-24"),
    ("CRM-02", "Migrate open opportunities", "Usman Tariq", "In Progress", "2026-10-31"),
    ("CRM-02", "Switch off legacy CRM", "Fatima Khan", "Not Started", "2026-11-30"),
]

RISKS = [
    # project, risk, likelihood (1-3), impact (1-3), mitigation, owner
    ("ERP-01", "UAT defects exceed fix capacity before go-live", 2, 3,
     "Freeze scope; add 2 developers to defect triage in Nov.", "Ali Raza"),
    ("ERP-01", "Key finance user unavailable during UAT", 2, 2,
     "Nominate a trained backup approver this week.", "Fatima Khan"),
    ("ERP-01", "Legacy data quality issues in item master", 3, 2,
     "Weekly data-quality checks; owners sign off each load.", "Sana Iqbal"),
    ("CRM-02", "Budget overrun — change requests not controlled", 3, 3,
     "All changes through change board; freeze non-critical requests.", "Fatima Khan"),
    ("CRM-02", "Training delay pushes go-live", 3, 3,
     "Run catch-up sessions; record video walkthroughs for absentees.", "Usman Tariq"),
    ("CRM-02", "Users revert to old CRM / spreadsheets", 2, 2,
     "Switch off legacy CRM on go-live; weekly adoption report.", "Fatima Khan"),
]

NEXT_STEPS = [
    # project, step, owner, due
    ("ERP-01", "Close UAT Round 1 — sign off Finance results", "Ali Raza", "2026-10-24"),
    ("ERP-01", "Start UAT Round 2 planning & test data refresh", "Sana Iqbal", "2026-10-27"),
    ("ERP-01", "Confirm hypercare roster for December", "Ali Raza", "2026-11-05"),
    ("CRM-02", "Reschedule Batch 2 training & confirm attendance", "Usman Tariq", "2026-10-20"),
    ("CRM-02", "Change-board review of pending change requests", "Fatima Khan", "2026-10-22"),
    ("CRM-02", "Draft go-live checklist & comms plan", "Fatima Khan", "2026-10-28"),
]


def sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.alignment = HDR_FILL, HDR_FONT, Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v)
    ws.freeze_panes = "A2"
    return ws


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Projects",
          ["ProjectID", "ProjectName", "Manager", "Sponsor", "StartDate", "EndDate",
           "BudgetPKR", "SpentPKR"],
          PROJECTS, [12, 26, 16, 20, 12, 12, 14, 14])
    sheet(wb, "Milestones",
          ["ProjectID", "Milestone", "DueDate", "Status", "Owner"],
          MILESTONES, [12, 46, 12, 14, 16])
    sheet(wb, "Tasks",
          ["ProjectID", "Task", "Owner", "Status", "DueDate"],
          TASKS, [12, 44, 16, 14, 12])
    sheet(wb, "Risks",
          ["ProjectID", "Risk", "Likelihood(1-3)", "Impact(1-3)", "Mitigation", "Owner"],
          RISKS, [12, 50, 16, 12, 52, 16])
    sheet(wb, "NextSteps",
          ["ProjectID", "NextStep", "Owner", "DueDate"],
          NEXT_STEPS, [12, 56, 16, 12])
    wb.save("project-data.xlsx")
    print("Wrote project-data.xlsx: 2 projects, "
          f"{len(MILESTONES)} milestones, {len(TASKS)} tasks, "
          f"{len(RISKS)} risks, {len(NEXT_STEPS)} next steps.")


if __name__ == "__main__":
    main()
