"""
PDF Status Report Generator — by Muhammad Bilal
=================================================
PROBLEM IT SOLVES
-----------------
Every week a PM rebuilds the same status report by hand: copy numbers from
a tracker into slides, recolour RAG cells, retype milestones and risks.
It takes 1–2 hours, the numbers drift between versions, and two projects
never look the same.

WHAT THIS DOES (100% local — no APIs, no network, no cloud)
------------------------------------------------------------
1. Reads project-data.xlsx (Projects / Milestones / Tasks / Risks / NextSteps).
2. Computes KPIs per project: task completion %, budget used %, RAG status,
   milestones on track, open risks.
3. Generates a polished multi-page PDF per project:
   cover page -> KPI summary -> milestone timeline -> risks & issues ->
   next steps. Headers, coloured tables, budget bar, page numbers.

USAGE:
    python generate_report.py                # all projects in project-data.xlsx
    python generate_report.py --project ERP-01
    python generate_report.py --input my-data.xlsx --output ./reports

Output: output/Status-Report-<ProjectName>-<YYYY-MM-DD>.pdf
"""
from __future__ import annotations

import argparse
import logging
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from fpdf import FPDF
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "project-data.xlsx"
DEFAULT_OUTPUT = BASE_DIR / "output"
NAVY = (14, 76, 129)          # headers / accents
DARK = (30, 41, 59)           # body text
GREY = (100, 116, 139)        # muted text
LIGHT = (238, 244, 250)       # card background
GREEN = (21, 128, 61)
AMBER = (180, 83, 9)
RED = (185, 28, 28)
GREEN_BG = (220, 252, 231)
AMBER_BG = (254, 243, 199)
RED_BG = (254, 226, 226)
ROW_ALT = (244, 247, 253)
WHITE = (255, 255, 255)

STATUS_FILL = {"Done": GREEN_BG, "In Progress": (219, 234, 254),
               "Delayed": RED_BG, "Not Started": (241, 245, 249)}
STATUS_TEXT = {"Done": GREEN, "In Progress": (30, 64, 175),
               "Delayed": RED, "Not Started": GREY}

logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s [%(levelname)s] %(message)s",
                    datefmt="%Y-%m-%d %H:%M:%S")
log = logging.getLogger("status-report")


def safe(text) -> str:
    """Map common Unicode punctuation to ASCII equivalents and drop anything
    outside latin-1, so core PDF fonts never crash on real PM data
    (em-dashes, arrows, curly quotes, bullets typed in Excel)."""
    if text is None:
        return ""
    table = {
        "—": " - ", "–": "-", "→": "->", "•": "-", "·": "-",
        "“": '"', "”": '"', "‘": "'", "’": "'",
        "…": "...", "×": "x", "✓": "[x]", "⚠": "!",
    }
    out = str(text)
    for k, v in table.items():
        out = out.replace(k, v)
    return out.encode("latin-1", errors="ignore").decode("latin-1")


# ---------------------------------------------------------------------------
# DATA MODEL
# ---------------------------------------------------------------------------
@dataclass
class Project:
    pid: str
    name: str
    manager: str
    sponsor: str
    start: str
    end: str
    budget: float
    spent: float
    milestones: list[dict] = field(default_factory=list)
    tasks: list[dict] = field(default_factory=list)
    risks: list[dict] = field(default_factory=list)
    next_steps: list[dict] = field(default_factory=list)

    @property
    def completion_pct(self) -> float:
        if not self.tasks:
            return 0.0
        done = sum(1 for t in self.tasks if t["status"] == "Done")
        return round(100.0 * done / len(self.tasks), 1)

    @property
    def budget_used_pct(self) -> float:
        return round(100.0 * self.spent / self.budget, 1) if self.budget else 0.0

    @property
    def rag(self) -> str:
        """Deterministic RAG: budget pressure and delayed milestones drive it."""
        if self.budget_used_pct > 95 or any(m["status"] == "Delayed" for m in self.milestones):
            return "RED"
        if self.budget_used_pct > 85 or any(
                t["status"] == "Delayed" for t in self.tasks):
            return "AMBER"
        return "GREEN"

    @property
    def milestones_on_track(self) -> str:
        good = sum(1 for m in self.milestones if m["status"] in ("Done", "In Progress"))
        return f"{good}/{len(self.milestones)}"


def load_data(xlsx: Path) -> list[Project]:
    if not xlsx.exists():
        raise FileNotFoundError(f"Input workbook not found: {xlsx}")
    wb = load_workbook(xlsx, data_only=True)
    required = {"Projects", "Milestones", "Tasks", "Risks", "NextSteps"}
    missing = required - set(wb.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheet(s): {sorted(missing)}")

    def rows(name: str) -> list[dict]:
        ws = wb[name]
        headers = [c.value for c in ws[1]]
        out = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in r):
                continue
            out.append(dict(zip(headers, r)))
        return out

    projects: dict[str, Project] = {}
    for r in rows("Projects"):
        projects[r["ProjectID"]] = Project(
            pid=r["ProjectID"], name=r["ProjectName"], manager=r["Manager"],
            sponsor=r["Sponsor"], start=str(r["StartDate"]), end=str(r["EndDate"]),
            budget=float(r["BudgetPKR"] or 0), spent=float(r["SpentPKR"] or 0))
    for r in rows("Milestones"):
        p = projects.get(r["ProjectID"])
        if p:
            p.milestones.append({"name": r["Milestone"], "due": str(r["DueDate"]),
                                 "status": r["Status"], "owner": r["Owner"]})
    for r in rows("Tasks"):
        p = projects.get(r["ProjectID"])
        if p:
            p.tasks.append({"name": r["Task"], "owner": r["Owner"],
                            "status": r["Status"], "due": str(r["DueDate"])})
    for r in rows("Risks"):
        p = projects.get(r["ProjectID"])
        if p:
            p.risks.append({"risk": r["Risk"], "likelihood": int(r["Likelihood(1-3)"] or 1),
                            "impact": int(r["Impact(1-3)"] or 1),
                            "mitigation": r["Mitigation"], "owner": r["Owner"]})
    for r in rows("NextSteps"):
        p = projects.get(r["ProjectID"])
        if p:
            p.next_steps.append({"step": r["NextStep"], "owner": r["Owner"],
                                 "due": str(r["DueDate"])})
    return list(projects.values())


# ---------------------------------------------------------------------------
# PDF BUILDER
# ---------------------------------------------------------------------------
class ReportPDF(FPDF):
    def __init__(self, project: Project, report_date: str):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.project = project
        self.report_date = report_date
        self.alias_nb_pages()
        self.set_auto_page_break(True, margin=22)
        self.set_margins(15, 15, 15)

    # -- bulletproof text handling -----------------------------------------
    # Core PDF fonts only support latin-1. Overriding cell/multi_cell means
    # NO string anywhere in this file (or in user data) can crash rendering.
    def cell(self, *args, **kwargs):
        # fpdf2 signature: cell(w, h, text, ...) — text is positional index 2
        if len(args) >= 3:
            args = args[:2] + (safe(args[2]),) + args[3:]
        elif "text" in kwargs:
            kwargs["text"] = safe(kwargs["text"])
        return super().cell(*args, **kwargs)

    def multi_cell(self, *args, **kwargs):
        # fpdf2 signature: multi_cell(w, h, text, ...) — text is positional index 2
        if len(args) >= 3:
            args = args[:2] + (safe(args[2]),) + args[3:]
        elif "text" in kwargs:
            kwargs["text"] = safe(kwargs["text"])
        return super().multi_cell(*args, **kwargs)

    # -- chrome ------------------------------------------------------------
    def footer(self):
        if self.page_no() == 1:
            return  # cover page stays clean
        self.set_y(-15)
        self.set_font("helvetica", "I", 8)
        self.set_text_color(*GREY)
        self.cell(0, 10, f"{safe(self.project.name)}  |  Status Report  |  "
                         f"Page {self.page_no()}/{{nb}}", align="C")

    def section_title(self, text: str):
        text = safe(text)
        self.set_font("helvetica", "B", 16)
        self.set_text_color(*NAVY)
        self.cell(0, 10, text, new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*NAVY)
        self.line(15, self.get_y(), 195, self.get_y())
        self.ln(4)

    def body(self, text: str, size: int = 10):
        text = safe(text)
        self.set_font("helvetica", "", size)
        self.set_text_color(*DARK)
        self.multi_cell(0, 5.5, text)

    # -- wrapping helper ---------------------------------------------------
    def wrap(self, text: str, width: float, size: int = 9) -> list[str]:
        """Split text into lines that fit `width` at `size` pt Helvetica."""
        self.set_font("helvetica", "", size)
        words, lines, cur = safe(text).split(), [], ""
        for w in words:
            trial = f"{cur} {w}".strip()
            if self.get_string_width(trial) <= width - 2:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        return lines or [""]

    def table_row(self, cells: list[tuple[str, float, tuple, tuple]],
                  line_h: float = 5.2, size: int = 9, bold: bool = False):
        """cells: (text, width, fill_rgb, text_rgb). Row height fits longest cell."""
        wrapped = [self.wrap(t, w, size) for t, w, _f, _c in cells]
        n_lines = max(len(w) for w in wrapped)
        h = n_lines * line_h + 2
        if self.get_y() + h > 275:
            self.add_page()
        x0, y0 = self.get_x(), self.get_y()
        for (text, w, fill, tcol), lines in zip(cells, wrapped):
            self.set_xy(x0, y0)
            self.set_fill_color(*fill)
            self.set_text_color(*tcol)
            self.set_font("helvetica", "B" if bold else "", size)
            self.cell(w, h, "", border=1, fill=True)
            self.set_xy(x0 + 1, y0 + 1)
            for ln in lines:
                self.cell(w - 2, line_h, ln, new_x="LMARGIN", new_y="NEXT")
                self.set_x(x0 + 1)
            x0 += w
        self.set_xy(15, y0 + h)

    # -- pages -------------------------------------------------------------
    def cover(self):
        self.add_page()
        self.ln(30)
        self.set_font("helvetica", "B", 13)
        self.set_text_color(*NAVY)
        self.cell(0, 8, "PROJECT STATUS REPORT", align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(4)
        self.set_font("helvetica", "B", 30)
        self.set_text_color(*DARK)
        self.multi_cell(0, 13, safe(self.project.name), align="C")
        self.ln(6)
        rag = self.project.rag
        fill = {"GREEN": GREEN_BG, "AMBER": AMBER_BG, "RED": RED_BG}[rag]
        txt = {"GREEN": GREEN, "AMBER": AMBER, "RED": RED}[rag]
        label = {"GREEN": "GREEN — ON TRACK", "AMBER": "AMBER — ATTENTION NEEDED",
                 "RED": "RED — ACTION REQUIRED"}[rag]
        self.set_fill_color(*fill)
        self.set_text_color(*txt)
        self.set_font("helvetica", "B", 14)
        x = 60
        self.set_x(x)
        self.cell(90, 14, safe(label), border=0, fill=True, align="C",
                  new_x="LMARGIN", new_y="NEXT")
        self.ln(10)
        self.set_font("helvetica", "", 12)
        self.set_text_color(*DARK)
        for label2, value in [
            ("Project Manager", self.project.manager),
            ("Sponsor", self.project.sponsor),
            ("Period", f"{self.project.start}  to  {self.project.end}"),
            ("Report date", self.report_date),
        ]:
            self.cell(45, 8, label2 + ":", align="R")
            self.set_font("helvetica", "B", 12)
            self.cell(0, 8, "  " + safe(value), new_x="LMARGIN", new_y="NEXT")
            self.set_font("helvetica", "", 12)
        self.ln(14)
        self.set_font("helvetica", "I", 10)
        self.set_text_color(*GREY)
        self.cell(0, 6, "Generated locally from project-data.xlsx — by Muhammad Bilal",
                  align="C")

    def kpi_page(self):
        self.add_page()
        self.section_title("1  ·  KPI Summary")
        p = self.project
        cards = [
            ("Task completion", f"{p.completion_pct}%", f"{sum(1 for t in p.tasks if t['status'] == 'Done')}/{len(p.tasks)} tasks done"),
            ("Budget used", f"{p.budget_used_pct}%", f"PKR {p.spent:,.0f} of PKR {p.budget:,.0f}"),
            ("Milestones on track", p.milestones_on_track, "Done + In Progress"),
            ("Open risks", str(len(p.risks)), "Tracked in Risks & Issues"),
        ]
        x0, y0 = 15, self.get_y()
        cw, ch, gap = 87, 30, 6
        for i, (label, value, sub) in enumerate(cards):
            x = x0 + (i % 2) * (cw + gap)
            y = y0 + (i // 2) * (ch + gap)
            self.set_xy(x, y)
            self.set_fill_color(*LIGHT)
            self.rect(x, y, cw, ch, style="F")
            self.set_xy(x + 4, y + 3)
            self.set_font("helvetica", "", 9)
            self.set_text_color(*GREY)
            self.cell(cw - 8, 5, label)
            self.set_xy(x + 4, y + 9)
            self.set_font("helvetica", "B", 18)
            self.set_text_color(*NAVY)
            self.cell(cw - 8, 10, value)
            self.set_xy(x + 4, y + 21)
            self.set_font("helvetica", "", 9)
            self.set_text_color(*DARK)
            self.cell(cw - 8, 5, sub)
        self.set_y(y0 + 2 * ch + 2 * gap + 6)
        # Budget bar
        self.set_font("helvetica", "B", 11)
        self.set_text_color(*DARK)
        self.cell(0, 7, "Budget consumption", new_x="LMARGIN", new_y="NEXT")
        bar_w, bar_h = 180, 8
        self.set_fill_color(226, 232, 240)
        self.rect(15, self.get_y(), bar_w, bar_h, style="F")
        pct = min(p.budget_used_pct, 100) / 100.0
        col = GREEN if p.budget_used_pct <= 85 else (AMBER if p.budget_used_pct <= 95 else RED)
        self.set_fill_color(*col)
        self.rect(15, self.get_y(), bar_w * pct, bar_h, style="F")
        self.ln(11)
        self.body(f"Budget used: {p.budget_used_pct}% (PKR {p.spent:,.0f} spent of "
                  f"PKR {p.budget:,.0f} approved). Remaining: PKR {p.budget - p.spent:,.0f}.")
        self.ln(2)
        self.body(f"Overall status: {p.rag} — " + {
            "GREEN": "the project is on track on budget and schedule.",
            "AMBER": "attention needed: budget pressure or a slipped task needs a decision this week.",
            "RED": "action required: budget overrun risk or a delayed milestone threatens the plan.",
        }[p.rag])

    def milestone_page(self):
        self.add_page()
        self.section_title("2  ·  Milestone Timeline")
        widths = [80, 25, 30, 45]
        headers = ["Milestone", "Due date", "Status", "Owner"]
        self.table_row([(h, w, NAVY, WHITE) for h, w in zip(headers, widths)], bold=True)
        for i, m in enumerate(self.project.milestones):
            fill = ROW_ALT if i % 2 else WHITE
            sfill = STATUS_FILL.get(m["status"], WHITE)
            scol = STATUS_TEXT.get(m["status"], DARK)
            self.table_row([
                (m["name"], widths[0], fill, DARK),
                (m["due"], widths[1], fill, DARK),
                (m["status"], widths[2], sfill, scol),
                (m["owner"], widths[3], fill, DARK),
            ])
        self.ln(4)
        self.body("Milestones are colour-coded: Done (green), In Progress (blue), "
                  "Delayed (red), Not Started (grey).")

    def risks_page(self):
        self.add_page()
        self.section_title("3  ·  Risks & Issues")
        widths = [55, 12, 12, 18, 55, 28]
        headers = ["Risk", "L", "I", "Rating", "Mitigation", "Owner"]
        self.table_row([(h, w, NAVY, WHITE) for h, w in zip(headers, widths)], bold=True)
        for i, r in enumerate(self.project.risks):
            score = r["likelihood"] * r["impact"]
            rating = f"{score}/9 ({'High' if score >= 6 else 'Medium' if score >= 3 else 'Low'})"
            rfill = RED_BG if score >= 6 else (AMBER_BG if score >= 3 else GREEN_BG)
            rcol = RED if score >= 6 else (AMBER if score >= 3 else GREEN)
            fill = ROW_ALT if i % 2 else WHITE
            self.table_row([
                (r["risk"], widths[0], fill, DARK),
                (str(r["likelihood"]), widths[1], fill, DARK),
                (str(r["impact"]), widths[2], fill, DARK),
                (rating, widths[3], rfill, rcol),
                (r["mitigation"], widths[4], fill, DARK),
                (r["owner"], widths[5], fill, DARK),
            ])
        self.ln(4)
        self.body("Rating = Likelihood (1–3) × Impact (1–3). High = 6–9, Medium = 3–5, Low = 1–2.")

    def next_steps_page(self):
        self.add_page()
        self.section_title("4  ·  Next Steps")
        for i, s in enumerate(self.project.next_steps, 1):
            self.set_font("helvetica", "B", 11)
            self.set_text_color(*NAVY)
            self.cell(10, 7, f"{i}.")
            self.set_font("helvetica", "", 11)
            self.set_text_color(*DARK)
            self.cell(0, 7, safe(s["step"]), new_x="LMARGIN", new_y="NEXT")
            self.set_font("helvetica", "", 9)
            self.set_text_color(*GREY)
            self.cell(10, 5, "")
            self.cell(0, 5, safe(f"Owner: {s['owner']}   -   Due: {s['due']}"),
                      new_x="LMARGIN", new_y="NEXT")
            self.ln(2)
        self.ln(6)
        self.set_font("helvetica", "I", 9)
        self.set_text_color(*GREY)
        self.multi_cell(0, 5, "This report was generated automatically from "
                              "project-data.xlsx. Figures are computed, not typed — "
                              "update the workbook and re-run to refresh the report.")


def sanitize(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-")


def build_report(project: Project, out_dir: Path, report_date: str) -> Path:
    pdf = ReportPDF(project, report_date)
    pdf.cover()
    pdf.kpi_page()
    pdf.milestone_page()
    pdf.risks_page()
    pdf.next_steps_page()
    out_dir.mkdir(parents=True, exist_ok=True)
    fname = f"Status-Report-{sanitize(project.name)}-{report_date}.pdf"
    out = out_dir / fname
    pdf.output(str(out))
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description="PDF Status Report Generator (100% local)")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT,
                        help="input workbook (default: project-data.xlsx)")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT,
                        help="output folder (default: ./output)")
    parser.add_argument("--project", type=str, default=None,
                        help="only generate for this ProjectID (default: all projects)")
    parser.add_argument("--date", type=str, default=str(date.today()),
                        help="report date YYYY-MM-DD (default: today)")
    args = parser.parse_args()
    try:
        projects = load_data(args.input)
    except (FileNotFoundError, ValueError) as exc:
        log.error("%s", exc)
        return 1
    if args.project:
        projects = [p for p in projects if p.pid == args.project]
        if not projects:
            log.error("No project with ID '%s' found in %s", args.project, args.input)
            return 1
    for p in projects:
        out = build_report(p, args.output, args.date)
        log.info("%s | RAG=%s | completion=%.1f%% | budget=%.1f%% -> %s",
                 p.name, p.rag, p.completion_pct, p.budget_used_pct, out.name)
    log.info("Done: %d report(s) in %s", len(projects), args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
