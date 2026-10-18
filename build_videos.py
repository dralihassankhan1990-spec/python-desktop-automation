"""
Builds the two demo videos for 04-Python-Desktop-Automation.

Every visual comes from REAL execution artefacts:
  - terminal cards: the actual watcher.log / generator console output
  - change-log table: the real change-log.csv
  - report shot: the real backup-report.html opened in Chromium
  - PDF pages: the real generated PDFs rendered with pdftoppm
  - title/problem/solution/result cards: rendered HTML cards
Narration: TTS via /opt/hatch/bin/tts (same tool as the Playwright demo video).

Outputs:
  folder-watcher-backup/folder-watcher-backup-demo.mp4
  pdf-status-report-generator/pdf-status-report-generator-demo.mp4
Usage: python build_videos.py [--watcher-only | --pdf-only]
"""
import csv
import html
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
WATCH = os.path.join(HERE, "folder-watcher-backup")
PDFGEN = os.path.join(HERE, "pdf-status-report-generator")
WORK = os.path.join(HERE, "_video_work")
AUD = os.path.join(WORK, "audio")
SEG = os.path.join(WORK, "segments")
VIS = os.path.join(WORK, "visuals")
for d in (AUD, SEG, VIS):
    os.makedirs(d, exist_ok=True)

TTS = "/opt/hatch/bin/tts"
FPS, VW, VH = 30, 1280, 720


def sh(*args, **kw):
    return subprocess.run(args, check=True, **kw)


def probe_dur(path):
    r = subprocess.run(["ffprobe", "-v", "error", "-show_entries",
                        "format=duration", "-of",
                        "default=noprint_wrappers=1:nokey=1", path],
                       capture_output=True, text=True, check=True)
    return float(r.stdout.strip())


def narrate(key, text):
    out = os.path.join(AUD, f"{key}.mp3")
    if not os.path.exists(out):
        sh(TTS, "speak", "--text-stdin", "--output", out,
           input=text.encode())
    return out


def seg_image(key, img, text, min_dur=2.0):
    """Landscape card/image with slow Ken Burns zoom + narration."""
    audio = narrate(key, text)
    dur = max(probe_dur(audio) + 0.8, min_dur)
    out = os.path.join(SEG, f"{key}.mp4")
    frames = int(dur * FPS)
    vf = (f"scale={VW}:{VH},zoompan=z='min(zoom+0.0012,1.12)':d={frames}:"
          f"x='iw/2-(iw/zoom/2)':y='ih/2-(ih/zoom/2)':s={VW}x{VH}:fps={FPS}")
    sh("ffmpeg", "-y", "-v", "error", "-loop", "1", "-i", img, "-i", audio,
       "-vf", vf, "-af", "apad", "-t", f"{dur:.2f}",
       "-c:v", "libx264", "-pix_fmt", "yuv420p", "-r", str(FPS),
       "-c:a", "aac", "-shortest", out)
    return out


def seg_image_fit(key, img, text, min_dur=2.0):
    """Portrait image (e.g. PDF page): fit on a dark 1280x720 canvas, no zoom."""
    audio = narrate(key, text)
    dur = max(probe_dur(audio) + 0.8, min_dur)
    out = os.path.join(SEG, f"{key}.mp4")
    vf = (f"scale=-2:660,pad={VW}:{VH}:(ow-iw)/2:(oh-ih)/2:color=0x0f172a,"
          f"fps={FPS}")
    sh("ffmpeg", "-y", "-v", "error", "-loop", "1", "-i", img, "-i", audio,
       "-vf", vf, "-af", "apad", "-t", f"{dur:.2f}",
       "-c:v", "libx264", "-pix_fmt", "yuv420p", "-r", str(FPS),
       "-c:a", "aac", "-shortest", out)
    return out


def concat(parts, final):
    lst = os.path.join(WORK, "concat.txt")
    with open(lst, "w") as f:
        for p in parts:
            f.write(f"file '{p}'\n")
    sh("ffmpeg", "-y", "-v", "error", "-f", "concat", "-safe", "0",
       "-i", lst, "-c", "copy", final)
    print(f"FINAL: {final} ({probe_dur(final):.1f}s)")


# ---------------------------------------------------------------------------
# Visual rendering (Playwright)
# ---------------------------------------------------------------------------
CARD_CSS = """
body { margin:0; font-family:'DejaVu Sans',Arial,sans-serif; }
.card { width:1280px; height:720px; display:flex; flex-direction:column;
        justify-content:center; align-items:center; text-align:center;
        background:linear-gradient(135deg,#0f172a 0%,#1e3a5f 55%,#0f766e 130%);
        color:#fff; position:relative; }
.kicker { font-size:26px; letter-spacing:6px; color:#5eead4; font-weight:bold;
          margin-bottom:18px; }
.title { font-size:60px; font-weight:bold; margin:0 80px 20px; line-height:1.15; }
.sub { font-size:25px; color:#cbd5e1; margin:0 110px; line-height:1.55; }
.footer { position:absolute; bottom:44px; font-size:19px; color:#94a3b8;
          letter-spacing:2px; }
"""

TERM_CSS = """
body { margin:0; background:#0b1220; font-family:'DejaVu Sans Mono',monospace; }
.term { padding:36px 44px; color:#e2e8f0; font-size:19px; line-height:1.65; }
.term .dim { color:#64748b; }
.term .grn { color:#4ade80; }
.term .amb { color:#fbbf24; }
.term .red { color:#f87171; }
.term .wht { color:#ffffff; font-weight:bold; }
.cap { color:#94a3b8; font-size:16px; margin-bottom:14px; letter-spacing:1px; }
"""

TABLE_CSS = """
body { margin:0; font-family:'DejaVu Sans',Arial,sans-serif; background:#0f172a; }
.sheet { margin:26px auto; width:1200px; background:#fff; border-radius:12px;
         overflow:hidden; box-shadow:0 12px 40px rgba(0,0,0,.45); }
.cap { background:#1f4e79; color:#fff; font-size:19px; font-weight:bold;
       padding:12px 20px; }
table { border-collapse:collapse; width:100%; font-size:14px; }
th { background:#d9e2f3; color:#1f4e79; padding:9px 11px; text-align:left;
     border:1px solid #b4c2dd; }
td { padding:9px 11px; border:1px solid #d5dbe8; color:#1e293b; }
tr:nth-child(even) td { background:#f4f7fd; }
"""


def card_html(kicker, title, sub, footer="PYTHON DESKTOP AUTOMATION  ·  BY MUHAMMAD BILAL"):
    return (f"<html><head><style>{CARD_CSS}</style></head><body><div class='card'>"
            f"<div class='kicker'>{kicker}</div><div class='title'>{title}</div>"
            f"<div class='sub'>{sub}</div><div class='footer'>{footer}</div>"
            f"</div></body></html>")


def term_html(caption, lines):
    body = "".join(f"<div>{ln}</div>" for ln in lines)
    return (f"<html><head><style>{TERM_CSS}</style></head><body>"
            f"<div class='term'><div class='cap'>{caption}</div>{body}</div>"
            f"</body></html>")


def table_html(caption, headers, rows):
    ths = "".join(f"<th>{html.escape(h)}</th>" for h in headers)
    trs = "".join("<tr>" + "".join(f"<td>{html.escape(str(v))}</td>" for v in r)
                  + "</tr>" for r in rows)
    return (f"<html><head><style>{TABLE_CSS}</style></head><body>"
            f"<div class='sheet'><div class='cap'>{caption}</div>"
            f"<table><tr>{ths}</tr>{trs}</table></div></body></html>")


def screenshot(html_content, out_path):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        pg = b.new_page(viewport={"width": VW, "height": VH})
        pg.set_content(html_content, wait_until="load")
        pg.wait_for_timeout(350)
        pg.screenshot(path=out_path)
        b.close()
    return out_path


def screenshot_file(file_url, out_path, wait=600):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        pg = b.new_page(viewport={"width": VW, "height": VH})
        pg.goto(file_url, wait_until="load")
        pg.wait_for_timeout(wait)
        pg.screenshot(path=out_path)
        b.close()
    return out_path


# ---------------------------------------------------------------------------
# VIDEO 1 — Folder Watcher
# ---------------------------------------------------------------------------
def build_watcher():
    V = lambda n: os.path.join(VIS, n)
    print("== building watcher visuals (real artefacts) ==")

    # Real watcher log lines (trimmed to the proof run, paths shortened)
    log_lines = []
    with open(os.path.join(WATCH, "watcher.log")) as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            line = line.replace(
                "/home/hatch/workspace/automation-build/04-Python-Desktop-Automation/"
                "folder-watcher-backup/", "")
            if "Pruned" in line:
                log_lines.append(f"<span class='amb'>{html.escape(line[20:])}</span>")
            elif "DELETED" in line:
                log_lines.append(f"<span class='red'>{html.escape(line[20:])}</span>")
            elif "CREATED" in line or "MODIFIED" in line:
                log_lines.append(f"<span class='grn'>{html.escape(line[20:])}</span>")
            else:
                log_lines.append(f"<span class='dim'>{html.escape(line[20:])}</span>")
    log_lines = log_lines[:11]  # proof run
    screenshot(term_html("TERMINAL — REAL RUN: python watcher.py --duration 45",
                         log_lines), V("w_term.png"))

    # Real backup tree
    tree = subprocess.run(["find", os.path.join(WATCH, "backup"), "-type", "f"],
                          capture_output=True, text=True, check=True).stdout
    tree_lines = [f"<span class='dim'>$ find backup -type f</span>"]
    for p in sorted(tree.strip().split("\n")):
        rel = p.replace(WATCH + "/", "")
        tree_lines.append(f"<span class='wht'>{html.escape(rel)}</span>")
    tree_lines.append("<span class='grn'>5 versioned backups — both edited versions kept</span>")
    screenshot(term_html("TERMINAL — REAL OUTPUT: versioned backup files",
                         tree_lines), V("w_tree.png"))

    # Real change-log.csv
    with open(os.path.join(WATCH, "change-log.csv"), newline="") as f:
        rows = list(csv.reader(f))
    screenshot(table_html("change-log.csv — the real event ledger",
                          rows[0], rows[1:]), V("w_csv.png"))

    # Real HTML report
    screenshot_file("file://" + os.path.join(WATCH, "backup-report.html"),
                    V("w_report.png"), wait=800)

    screenshot(card_html("PYTHON DESKTOP AUTOMATION · 1 OF 2",
                         "Folder Watcher &<br>Auto-Backup",
                         "Never lose a file version again."),
               V("w_title.png"))
    screenshot(card_html("THE PROBLEM",
                         "One overwrite destroys<br>the previous version",
                         "A project team shares one folder. Someone overwrites the "
                         "status tracker. A file gets deleted by accident. "
                         "Cloud sync keeps only the latest copy — so the old version "
                         "is gone forever."),
               V("w_problem.png"))
    screenshot(card_html("THE SOLUTION",
                         "Watch. Version. Log. Report.",
                         "The watcher polls the folder every few seconds. When a file "
                         "is created or changed, it saves a timestamped copy, keeps "
                         "the last five versions of every file, and logs every event. "
                         "Everything runs locally. No cloud, no API, no subscription."),
               V("w_solution.png"))
    screenshot(card_html("THE RESULT",
                         "Six events. Zero lost versions.",
                         "Five timestamped backups, a complete change log, a daily "
                         "report — and a retention test that collapsed seven edits "
                         "down to the five newest versions. Set it and forget it."),
               V("w_result.png"))

    print("== assembling watcher video ==")
    parts = [
        seg_image("w_title", V("w_title.png"),
                  "Folder Watcher and Auto Backup. Never lose a file version again.", 4.0),
        seg_image("w_problem", V("w_problem.png"),
                  "Here is the problem. A project team shares one folder. Someone "
                  "overwrites the status tracker. A deliverable gets deleted by "
                  "accident. Cloud sync keeps only the latest copy, so one overwrite "
                  "destroys the previous version forever. Recovery becomes a hunt "
                  "through chat history and email.", 5.0),
        seg_image("w_solution", V("w_solution.png"),
                  "The solution is this watcher. It polls the folder every few "
                  "seconds. When a file is created or changed, it saves a timestamped "
                  "copy, keeps the last five versions of every file, and logs every "
                  "event to a CSV file. At the end of the run it builds an HTML "
                  "report. Everything runs locally on your machine. No cloud, no API, "
                  "no subscription.", 6.0),
        seg_image("w_term", V("w_term.png"),
                  "Watch it work for real. I start the watcher, then run a demo that "
                  "creates a file, edits the project plan twice, adds a spreadsheet, "
                  "and deletes a scratch file. The watcher catches all six events live.", 6.0),
        seg_image("w_tree", V("w_tree.png"),
                  "Here is the proof. Five timestamped backup copies — including both "
                  "versions of the edited plan. Nothing is overwritten, ever.", 5.0),
        seg_image("w_csv", V("w_csv.png"),
                  "And the change log records every event: created, modified, "
                  "deleted — with the exact backup path for each one.", 5.0),
        seg_image("w_report", V("w_report.png"),
                  "The daily HTML report summarizes it all: events today, versions "
                  "per file, and the recent history.", 5.0),
        seg_image("w_result", V("w_result.png"),
                  "Six events handled. Zero lost versions. The retention test proved "
                  "pruning too: seven edits collapsed to the five newest versions. "
                  "Set it and forget it. Your versions are safe.", 6.0),
    ]
    concat(parts, os.path.join(WATCH, "folder-watcher-backup-demo.mp4"))


# ---------------------------------------------------------------------------
# VIDEO 2 — PDF Status Report Generator
# ---------------------------------------------------------------------------
def build_pdf():
    V = lambda n: os.path.join(VIS, n)
    print("== building PDF-generator visuals (real artefacts) ==")

    # Fresh real run, capture console output (prefer an interpreter with fpdf2)
    gen = None
    for interp in (sys.executable, "/tmp/venv-auto/bin/python"):
        if not interp or not os.path.exists(interp):
            continue
        gen = subprocess.run([interp, "generate_report.py"], cwd=PDFGEN,
                             capture_output=True, text=True)
        if gen.returncode == 0:
            break
    if gen is None or gen.returncode != 0:
        raise RuntimeError("generate_report.py failed: "
                           + (gen.stderr if gen else "no interpreter found"))
    console = [html.escape(l) for l in gen.stderr.strip().split("\n") if l.strip()]
    term_lines = [f"<span class='dim'>$ python generate_report.py</span>"]
    for ln in console:
        if "RAG=GREEN" in ln:
            term_lines.append(f"<span class='grn'>{ln}</span>")
        elif "RAG=RED" in ln:
            term_lines.append(f"<span class='red'>{ln}</span>")
        else:
            term_lines.append(f"<span class='wht'>{ln}</span>")
    screenshot(term_html("TERMINAL — REAL RUN: generate_report.py", term_lines),
               V("p_term.png"))

    # Real workbook — Projects sheet
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(PDFGEN, "project-data.xlsx"))
    ws = wb["Projects"]
    headers = [c.value for c in ws[1]]
    prows = [[c for c in r] for r in ws.iter_rows(min_row=2, values_only=True)
             if r[0]]
    prows = [[str(v) for v in r] for r in prows]
    screenshot(table_html("project-data.xlsx — Projects sheet (the only input)",
                          headers, prows), V("p_xlsx.png"))

    # Real PDF pages (ERP cover + KPI, CRM risks) rendered at 110 dpi
    for tag, pdf in [("erp", "Status-Report-Riverdale-ERP-Rollout-2026-10-16.pdf"),
                     ("crm", "Status-Report-Harbor-CRM-Upgrade-2026-10-16.pdf")]:
        subprocess.run(["pdftoppm", "-png", "-r", "110", "-f", "1", "-l", "5",
                        os.path.join(PDFGEN, "output", pdf),
                        os.path.join(VIS, f"p_{tag}")],
                       check=True, capture_output=True)

    screenshot(card_html("PYTHON DESKTOP AUTOMATION · 2 OF 2",
                         "PDF Status Report<br>Generator",
                         "Weekly status reports in seconds, not hours."),
               V("p_title.png"))
    screenshot(card_html("THE PROBLEM",
                         "Two hours a week,<br>rebuilt by hand",
                         "Every week, a project manager rebuilds the same status "
                         "report by hand. Copy numbers into slides. Recolour the "
                         "RAG cells. Retype the milestones. Two hours gone, numbers "
                         "drift between versions, and no two projects ever look "
                         "the same."),
               V("p_problem.png"))
    screenshot(card_html("THE SOLUTION",
                         "Excel in. Polished PDF out.",
                         "The generator reads one Excel workbook with your projects, "
                         "milestones, tasks, risks and next steps. It computes "
                         "completion, budget used and the RAG status automatically. "
                         "Then it builds a five page PDF per project: cover, KPI "
                         "summary, milestone timeline, risks, and next steps."),
               V("p_solution.png"))
    screenshot(card_html("THE RESULT",
                         "Same input. Same output. Every week.",
                         "Update the workbook, re-run, and the report is ready to "
                         "send. That is two hours back, every single week."),
               V("p_result.png"))

    print("== assembling PDF-generator video ==")
    parts = [
        seg_image("p_title", V("p_title.png"),
                  "PDF Status Report Generator. Weekly reports in seconds, not hours.", 4.0),
        seg_image("p_problem", V("p_problem.png"),
                  "Every week, a project manager rebuilds the same status report by "
                  "hand. Copy numbers into slides. Recolour the RAG cells. Retype "
                  "the milestones. Two hours gone, numbers drift between versions, "
                  "and no two projects ever look the same.", 5.0),
        seg_image("p_solution", V("p_solution.png"),
                  "This generator reads one Excel workbook with your projects, "
                  "milestones, tasks, risks and next steps. It computes completion, "
                  "budget used and the RAG status automatically. Then it builds a "
                  "polished five page PDF per project: cover, KPI summary, milestone "
                  "timeline, risks, and next steps.", 6.0),
        seg_image("p_xlsx", V("p_xlsx.png"),
                  "The only input is this workbook. Two fictional projects, with "
                  "their milestones, tasks, risks and next steps.", 5.0),
        seg_image("p_term", V("p_term.png"),
                  "Here it is running for real. Two projects in, two PDFs out. "
                  "Riverdale ERP: green, fifty percent complete, sixty two percent "
                  "of budget. Harbor CRM: red, because training slipped.", 6.0),
        seg_image_fit("p_cover", V("p_erp-1.png"),
                      "The cover carries the project name and a RAG badge. Riverdale "
                      "is green, on track.", 5.0),
        seg_image_fit("p_kpi", V("p_erp-2.png"),
                      "The KPI page shows four cards — completion, budget used, "
                      "milestones on track, open risks — plus a budget bar.", 5.0),
        seg_image_fit("p_risks", V("p_crm-4.png"),
                      "Risks carry a likelihood times impact rating, colour coded "
                      "high, medium and low, with the mitigation and the owner.", 5.0),
        seg_image("p_result", V("p_result.png"),
                  "Same input, same output, every week. Update the workbook, re-run, "
                  "and the report is ready to send. That is two hours back, every "
                  "single week.", 6.0),
    ]
    concat(parts, os.path.join(PDFGEN, "pdf-status-report-generator-demo.mp4"))


if __name__ == "__main__":
    only = sys.argv[1] if len(sys.argv) > 1 else "all"
    if only in ("all", "--watcher-only"):
        build_watcher()
    if only in ("all", "--pdf-only"):
        build_pdf()
    print("Done.")
